"""Generate the sample DOCX files used by the merge demo.

Usage:
    python samples/generate_samples.py

Requires: python-docx, Pillow.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"

BRAND = RGBColor(0x24, 0x57, 0xD8)
DEEP = RGBColor(0x1C, 0x46, 0xAD)
GREY = RGBColor(0x5B, 0x64, 0x74)


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = tc_pr.makeelement(qn("w:shd"), {})
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_text(cell, text: str, *, bold: bool = False, color: RGBColor | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    if color is not None:
        run.font.color.rgb = color


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        set_cell_shading(table.rows[0].cells[index], "2457D8")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
    document.add_paragraph()


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_numbers(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Number")


def build_logo() -> Path:
    ASSETS.mkdir(parents=True, exist_ok=True)
    path = ASSETS / "brand-logo.png"
    image = Image.new("RGB", (480, 240), "white")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((24, 24, 456, 216), radius=28, fill="#2457d8")
    draw.ellipse((52, 62, 176, 186), fill="#ffffff")
    draw.rectangle((150, 96, 420, 152), fill="#eef3ff")
    draw.rectangle((186, 112, 360, 136), fill="#2457d8")
    image.save(path)
    return path


def configure_document(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        heading = document.styles[name]
        heading.font.name = "微软雅黑"
        heading.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        heading.font.color.rgb = DEEP
    document.styles["Title"].font.color.rgb = BRAND


def build_main() -> Path:
    document = Document()
    configure_document(document)

    title = document.add_paragraph("OpenHarness 交付文档", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("主文档 + 副文档自动合并示例")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.color.rgb = GREY
    subtitle.runs[0].font.size = Pt(10)

    document.add_paragraph("本文档由主文档和多个副文档自动合并生成。目录、标题编号与页码可在 Word 中按 F9 更新。")
    document.add_paragraph("${toc}")

    document.add_heading("项目概述", level=1)
    document.add_paragraph(
        "OpenHarness 是一套面向研发团队的智能体编排平台。本交付文档汇总产品能力与实施计划，"
        "用于演示按标记提取副文档内容并合并到主文档的完整流程。"
    )
    add_bullets(
        document,
        [
            "主文档提供封面、目录、总体说明与附录。",
            "副文档只贡献被 tag-start / tag-end 标记包裹的内容。",
            "合并结果保留原有段落、表格、图片与样式。",
        ],
    )

    document.add_paragraph("${merge-point}")

    document.add_heading("附录：使用说明", level=1)
    document.add_paragraph(
        "在 Word 中打开合并结果后，按 Ctrl+A 再按 F9 即可刷新目录页码。"
        "如果希望重新生成，请回到上传页面替换文件后再次合并。"
    )

    target = ROOT / "主文档.docx"
    document.save(target)
    return target


def build_product(logo: Path) -> Path:
    document = Document()
    configure_document(document)

    document.add_heading("产品介绍", level=1)
    document.add_paragraph("${product-start}")

    document.add_heading("核心能力", level=2)
    document.add_paragraph(
        "平台以任务为中心组织模型、工具与知识库，支持可视化编排、运行追踪和结果沉淀。"
    )
    document.add_paragraph("下图展示产品标识：")
    document.add_picture(str(logo), width=Inches(2.6))

    add_table(
        document,
        ["能力模块", "说明", "交付状态"],
        [
            ["任务编排", "拖拽式流程设计，支持条件分支与并行执行", "已交付"],
            ["工具市场", "内置文档、检索、代码执行等工具", "已交付"],
            ["运行观测", "调用链、Token 与耗时统计", "灰度中"],
        ],
    )

    document.add_heading("适用场景", level=2)
    add_bullets(
        document,
        [
            "研发团队的知识检索与文档问答。",
            "跨系统数据整理与定期报告生成。",
            "面向业务人员的低代码智能助手搭建。",
        ],
    )
    document.add_paragraph("${product-end}")
    document.add_paragraph("这段位于标记之外，不应出现在合并结果中。")

    target = ROOT / "副文档-产品介绍.docx"
    document.save(target)
    return target


def build_plan() -> Path:
    document = Document()
    configure_document(document)

    document.add_heading("实施计划", level=1)
    document.add_paragraph("${plan-start}")

    document.add_heading("阶段划分", level=2)
    add_numbers(
        document,
        [
            "第 1 周：环境准备、账号与权限梳理。",
            "第 2-3 周：主数据接入与知识库初始化。",
            "第 4 周：试点团队培训与流程试运行。",
            "第 5 周：验收、问题收敛与正式上线。",
        ],
    )

    add_table(
        document,
        ["阶段", "关键交付物", "负责人"],
        [
            ["准备", "环境清单、权限矩阵", "平台组"],
            ["接入", "数据源清单、知识库索引", "数据组"],
            ["试点", "试运行报告、问题清单", "业务组"],
            ["上线", "验收报告、运维手册", "项目组"],
        ],
    )

    document.add_heading("风险与应对", level=2)
    add_bullets(
        document,
        [
            "历史文档格式不统一：先做抽样清洗，再批量导入。",
            "权限边界不清：上线前完成角色矩阵评审。",
            "试点反馈集中：预留一周缓冲期处理问题。",
        ],
    )
    document.add_paragraph("${plan-end}")
    document.add_paragraph("这段位于标记之外，不应出现在合并结果中。")

    target = ROOT / "副文档-实施计划.docx"
    document.save(target)
    return target


def build_invalid() -> Path:
    document = Document()
    configure_document(document)
    document.add_heading("异常标记示例", level=1)
    document.add_paragraph("${broken-start}")
    document.add_paragraph("这个文档只有开始标记，没有结束标记，用于演示警告提示。")
    document.add_paragraph("${orphan-end}")
    target = ROOT / "副文档-异常标记.docx"
    document.save(target)
    return target


def main() -> None:
    logo = build_logo()
    created = [
        build_main(),
        build_product(logo),
        build_plan(),
        build_invalid(),
    ]
    for path in created:
        print(f"created: {path.relative_to(ROOT.parent)}")


if __name__ == "__main__":
    main()
