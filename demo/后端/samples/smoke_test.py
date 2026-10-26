"""End-to-end smoke test for the DOCX merge service.

Start the backend first:
    java -jar target/doc-merge-demo-1.0.0.jar

Then run:
    python samples/smoke_test.py
"""

from __future__ import annotations

import io
import os
import sys
import zipfile
from pathlib import Path

import requests
from docx import Document

BASE_URL = os.environ.get("BASE_URL", "http://localhost:8080")
ROOT = Path(__file__).resolve().parent


def fail(message: str) -> None:
    print(f"FAIL: {message}")
    raise SystemExit(1)


def check(condition: bool, message: str) -> None:
    if not condition:
        fail(message)
    print(f"  ok - {message}")


def upload_inspect(path: Path) -> dict:
    with path.open("rb") as handle:
        response = requests.post(
            f"{BASE_URL}/api/documents/inspect",
            files={"file": (path.name, handle, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            timeout=60,
        )
    payload = response.json()
    check(response.ok and payload.get("code") == 0, f"inspect {path.name} returns code 0")
    return payload["data"]


def main() -> None:
    main_file = ROOT / "主文档.docx"
    scope = ROOT / "检测范围.docx"
    plan = ROOT / "实施计划.docx"
    invalid = ROOT / "异常标记.docx"
    for path in (main_file, scope, plan, invalid):
        if not path.exists():
            fail(f"missing sample: {path} (run samples/generate_samples.py)")

    print("[1] inspect samples")
    main_info = upload_inspect(main_file)
    check(main_info["tagCount"] == 0, "main document has no name-start/name-end tags")
    scope_info = upload_inspect(scope)
    check(scope_info["tagCount"] == 2, "scope document has 2 tag occurrences")
    check(any(tag["status"] == "PAIRED" for tag in scope_info["tags"]), "scope tag is paired")
    invalid_info = upload_inspect(invalid)
    check(len(invalid_info.get("warnings") or []) > 0, "invalid document reports warnings")

    print("[2] merge two sub documents")
    files = [("mainFile", (main_file.name, main_file.open("rb"), "application/octet-stream"))]
    files.append(("subFiles", (scope.name, scope.open("rb"), "application/octet-stream")))
    files.append(("subFiles", (plan.name, plan.open("rb"), "application/octet-stream")))
    response = requests.post(
        f"{BASE_URL}/api/documents/merge",
        files=files,
        data={"outputName": "合并结果"},
        timeout=120,
    )
    payload = response.json()
    check(response.ok and payload.get("code") == 0, "merge returns code 0")
    result = payload["data"]
    check(result["fileName"] == "合并结果.docx", "output name is normalized")
    check(result["size"] > 0, "output size is positive")
    check(any(tag["status"] == "PAIRED" for tag in result.get("tags") or []), "merged tags are paired")
    check(not (result.get("warnings") or []), "valid merge has no warnings")

    print("[3] download and validate output")
    download = requests.get(f"{BASE_URL}{result['downloadUrl']}", timeout=60)
    check(download.ok and len(download.content) > 0, "download returns bytes")
    check(
        "wordprocessingml.document" in download.headers.get("Content-Type", ""),
        "download content type is DOCX",
    )

    target = ROOT / "合并结果.docx"
    target.write_bytes(download.content)

    document = Document(io.BytesIO(download.content))
    text = "\n".join(p.text for p in document.paragraphs)
    table_text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    combined = text + "\n" + table_text
    check("检测项目与记录要求" in combined, "scope content is present")
    check("时间安排" in combined, "plan content is present")
    check("电压、电流、温升和绝缘状态" in combined, "table content is preserved")
    check("附录：记录说明" in combined, "main document appendix is preserved")
    check("这段说明在标记之外" not in combined, "content outside markers is excluded")
    check("${" not in combined, "markers are removed from output")

    with zipfile.ZipFile(io.BytesIO(download.content)) as archive:
        names = archive.namelist()
        document_xml = archive.read("word/document.xml").decode("utf-8")
        settings_xml = archive.read("word/settings.xml").decode("utf-8")
        check(any(name.startswith("word/media/") for name in names), "image part is copied")
        check("TOC \\o" in document_xml, "native TOC field is present")
        check("updateFields" in settings_xml, "updateFields is written")

    print(f"\nAll checks passed. Output saved to {target}")


if __name__ == "__main__":
    try:
        main()
    except requests.RequestException as error:
        fail(f"cannot reach backend: {error}")
